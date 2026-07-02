"""Import a completed hand-grading workbook pass.

Validates every row (score present, within [0, max_points]) and writes
``hand_grades_pass<N>.json`` keyed by artifact_id for the metrics pipeline.

Usage:
    uv run python scripts/import_hand_grades.py tests/eval/gold_set/hand_grading_pass1.xlsx --pass-num 1
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook")
    parser.add_argument("--pass-num", type=int, required=True, choices=(1, 2))
    args = parser.parse_args()

    ws = load_workbook(args.workbook, read_only=True).active
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    idx = {name: header.index(name) for name in ("artifact_id", "question", "max_points", "score", "notes")}

    grades: dict[str, dict] = {}
    problems: list[str] = []
    for n, row in enumerate(ws.iter_rows(min_row=2), start=2):
        artifact_id = row[idx["artifact_id"]].value
        if artifact_id is None:
            continue
        score = row[idx["score"]].value
        max_points = float(row[idx["max_points"]].value or 0)
        if score is None:
            problems.append(f"row {n}: missing score")
            continue
        try:
            score = float(score)
        except (TypeError, ValueError):
            problems.append(f"row {n}: non-numeric score {score!r}")
            continue
        if not 0 <= score <= max_points:
            problems.append(f"row {n}: score {score} outside [0, {max_points}]")
            continue
        grades[str(artifact_id)] = {
            "score": score, "max_points": max_points,
            "question": row[idx["question"]].value,
            "notes": row[idx["notes"]].value,
        }

    if problems:
        print(f"{len(problems)} problem rows (not imported):")
        for p in problems[:20]:
            print(" ", p)
    out = Path(args.workbook).parent / f"hand_grades_pass{args.pass_num}.json"
    out.write_text(json.dumps(grades, indent=2))
    print(f"imported {len(grades)} grades -> {out}")
    return 0 if not problems else 1


if __name__ == "__main__":
    raise SystemExit(main())
