"""Export the blinded hand-grading workbook for one pass.

One row per (submission, question) with the normalized answer text — machine
scores, tiers, and confidences are deliberately withheld. Row order is
shuffled per pass so pass 2 does not replay pass 1's sequence. Scores go in
the `score` column; the importer reads them back by the stable artifact id.

Usage:
    uv run python scripts/export_grading_workbook.py --pass-num 1
"""

from __future__ import annotations

import argparse
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from autograde_core import storage
from autograde_core.hashing import assignment_id_for

ASSIGNMENTS = ["gold-theory", "gold-algo"]


def rows_for(assignment: str) -> list[dict]:
    from sqlalchemy import select

    from autograde_core.models import Rubric
    from autograde_db.models import Artifact, Assignment, Submission
    from autograde_db.session import session_scope

    auuid = assignment_id_for(assignment)
    with session_scope() as s:
        rubric = Rubric.model_validate(s.get(Assignment, auuid).rubric)
        weights = {q.question_id: q.weight for q in rubric.questions}
        prompts = {q.question_id: q.prompt for q in rubric.questions}
        criteria = {
            q.question_id: "; ".join(f"{c.criterion} ({c.points})" for c in q.rubric_criteria)
            for q in rubric.questions
        }
        arts = s.execute(
            select(Artifact, Submission.student_hash)
            .join(Submission, Artifact.submission_id == Submission.id)
            .where(Submission.assignment_id == auuid)
            .where(Artifact.question_id.is_not(None))
        ).all()

    out = []
    for artifact, _hash in arts:
        content = storage.get_json("normalized-artifacts", artifact.content_ref)
        body = content.get("content", {})
        answer = (body.get("text") or body.get("source") or body.get("raw_text")
                  or f"[figure: OCR text] {body.get('ocr_text') or '(none)'}")
        qid = artifact.question_id
        out.append({
            "artifact_id": str(artifact.id), "assignment": assignment,
            "question": qid, "prompt": prompts.get(qid, ""),
            "criteria": criteria.get(qid, ""), "max_points": weights.get(qid, 0),
            "answer": str(answer)[:6000],
        })
    return out


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--pass-num", type=int, required=True, choices=(1, 2))
    parser.add_argument("--out", default="tests/eval/gold_set")
    args = parser.parse_args()

    rows: list[dict] = []
    for assignment in ASSIGNMENTS:
        rows.extend(rows_for(assignment))
    random.Random(1000 + args.pass_num).shuffle(rows)  # different order per pass

    wb = Workbook()
    ws = wb.active
    ws.title = f"pass {args.pass_num}"
    headers = ["artifact_id", "assignment", "question", "prompt", "criteria",
               "max_points", "answer", "score", "notes"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r[h] for h in headers[:7]] + [None, None])
    for col, width in zip(ws.columns, (8, 12, 9, 40, 40, 10, 90, 8, 30), strict=False):
        ws.column_dimensions[col[0].column_letter].width = width
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    path = Path(args.out) / f"hand_grading_pass{args.pass_num}.xlsx"
    wb.save(path)
    print(f"wrote {path} with {len(rows)} rows (blinded, shuffled for pass {args.pass_num})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
